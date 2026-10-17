"""Construye ASTROS_Campanas.xlsx: briefs, calendario semanal de contenido, guiones de WhatsApp y presupuesto.
Fechas de examen tomadas de admisiones/ASTROS_calendario_admision_2027.xlsx; precios de precios/tarifario.csv.
Uso: python3 construir_campanas.py [salida.xlsx]"""
import csv, datetime as dt, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAIZ = os.path.dirname(BASE)
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "ASTROS_Campanas.xlsx")
D = dt.date
PRECIO = {r["id"]: r for r in csv.DictReader(open(os.path.join(RAIZ, "precios", "tarifario.csv"), encoding="utf-8")) if r["activo"] == "si"}
def soles(pid): return f"S/ {float(PRECIO[pid]['precio_virtual']):,.0f}"

# código, nombre, universidad, examen (fecha), cierre inscripción, inicio, fin, público, oferta (ids tarifario), gancho, canales, ads S/ por día, días de ads, meta leads, meta inscritos
CAMPANAS = [
 ("SPRINT-ULIMA", "Sprint ULIMA", "ULIMA", D(2026, 11, 15), D(2026, 11, 6), D(2026, 10, 13), D(2026, 11, 15), "5.° de secundaria que dará el Examen de Escolares ULIMA",
  ["pre-pack12", "pre-mensual"], "Clase diagnóstica gratis de 45 min + plan de 4 semanas hasta el 15/11", "IG Reels, historias, WhatsApp, web", 20, 21, 40, 5),
 ("SPRINT-PUCP", "Sprint Primera Opción PUCP", "PUCP", D(2026, 11, 29), None, D(2026, 10, 14), D(2026, 11, 29), "5.° de secundaria que dará La Primera Opción PUCP (y ARS 3.ª convocatoria)",
  ["pre-pack12", "pre-mensual"], "Simulacro diagnóstico gratis tipo PUCP con solucionario + clase diagnóstica de 45 min", "IG, FB, grupos de padres, WhatsApp, web", 25, 28, 60, 8),
 ("REFUERZO-CEPU", "Refuerzo CEPU Cayetano", "UPCH", D(2026, 11, 29), None, D(2026, 10, 20), D(2026, 11, 29), "Alumnos de CEPU II (Pre Cayetano) antes de su 3.er examen",
  ["pre-hora", "pre-g3"], "Tutoría por curso (Biología, Química, RM) en hora suelta o en grupo de 2-3", "IG, referidos, WhatsApp", 10, 14, 15, 3),
 ("MEDICINA-UPCH", "Ciclo Medicina Cayetano", "UPCH", D(2027, 1, 17), D(2027, 1, 11), D(2026, 11, 16), D(2027, 1, 17), "5.° de secundaria y egresados que postulan a Medicina UPCH",
  ["pre-mensual"], "Webinar gratis 'Medicina UPCH en 8 semanas' + simulacro semanal", "IG Ads segmentado, FB, web, WhatsApp", 25, 35, 40, 6),
 ("VERANO-PUCP", "Ciclo Verano Talento PUCP", "PUCP", D(2027, 2, 14), None, D(2026, 12, 1), D(2027, 2, 14), "Egresados y quienes no ingresaron el 29/11 (Evaluación del Talento, fecha por confirmar)",
  ["pre-mensual", "pre-g3"], "'Segunda oportunidad en febrero': diagnóstico gratis + ciclo de verano", "IG Ads, WhatsApp a leads no ingresantes", 20, 30, 50, 8),
 ("VERANO-ULIMA", "Ciclo Verano ULIMA egresados", "ULIMA", D(2027, 3, 7), None, D(2026, 12, 1), D(2027, 3, 7), "Egresados que darán el examen ULIMA 2027-1B (fecha no oficial)",
  ["pre-mensual", "pre-pack12"], "Simulacro diagnóstico + ciclo de verano", "IG, FB, WhatsApp", 15, 25, 30, 5),
 ("CACHIMBOS", "Tutoría Cachimbos", "Universitario", D(2027, 3, 29), None, D(2026, 12, 1), D(2027, 4, 30), "Ingresantes 2027-1 de PUCP, ULIMA, UPC y UPCH (Cálculo, Física, Química, Estadística)",
  ["uni-practica", "uni-medio", "uni-ciclo", "uni-g3"], "Lista de espera en diciembre; paquetes por curso desde marzo", "IG, grupos de ingresantes, referidos", 15, 30, 40, 10),
 ("REFERIDOS", "Trae a tu tripulación", "Todas", None, None, D(2026, 11, 30), D(2027, 4, 30), "Alumnos y ex alumnos de ASTROS",
  [], "1 hora gratis por cada referido que compra un paquete (política de precios)", "WhatsApp, IG", 0, 0, 20, 5),
 ("ORGANICO", "Contenido de marca", "Todas", None, None, D(2026, 10, 13), D(2027, 4, 30), "Comunidad de @1.masesores",
  [], "Frases ASTROS, tips de estudio, testimonios", "IG, FB, TikTok", 0, 0, 0, 0),
]
CODIGOS = [c[0] for c in CAMPANAS]

NAVY, GOLD, TEAL, RED = "0B2533", "F9C74F", "1FA2A6", "E63946"
hf, hfill = Font(bold=True, color=NAVY), PatternFill("solid", fgColor=GOLD)
wrap = Alignment(wrap_text=True, vertical="top")
wb = Workbook()

def hoja(nombre, cols, widths, tab=TEAL):
    ws = wb.create_sheet(nombre); ws.sheet_properties.tabColor = tab
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        ws.cell(1, j, c).font = hf; ws.cell(1, j).fill = hfill; ws.cell(1, j).alignment = wrap
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"; ws.row_dimensions[1].height = 30
    return ws

# ---- Inicio
ini = wb.active; ini.title = "Inicio"; ini.sheet_properties.tabColor = NAVY
ini.column_dimensions["A"].width = 22; ini.column_dimensions["B"].width = 95
ini["A1"] = "ASTROS · Plan de campañas oct 2026 – abr 2027"; ini["A1"].font = Font(bold=True, size=16, color="FFFFFF")
ini["A1"].fill = ini["B1"].fill = PatternFill("solid", fgColor=NAVY)
for i, (a, b) in enumerate([
    ("Qué es", "El plan ejecutable de marketing: qué publicar cada semana, con qué texto, para qué campaña y con cuánto presupuesto. Se apoya en el cronograma de admisión (fechas y metas) y en el tarifario aprobado."),
    ("Campañas", "Una fila por campaña: objetivo, público, oferta con precios, gancho, canales, presupuesto y metas."),
    ("Calendario", "Semana a semana: 3 publicaciones + 1 difusión por WhatsApp. Cambia la columna Estado a 'Publicado' cuando lo subas."),
    ("WhatsApp", "Guiones listos para copiar: primer contacto, seguimiento, cierre, recordatorio de inscripción, post-resultado y referidos."),
    ("Presupuesto", "Inversión en anuncios por campaña y costo por alumno inscrito cuando cargues los resultados reales."),
    ("Medición", "Cada lead entra al SIG con su código de campaña (SPRINT-PUCP, etc.). El Dashboard del SIG muestra leads, inscritos y ventas por campaña."),
    ("Ojo", "Las fechas de PUCP Talento (~14/02) y ULIMA egresados (~07/03) no son oficiales todavía: revisa a inicios de noviembre y de enero."),
], 3):
    ini[f"A{i}"] = a; ini[f"A{i}"].font = Font(bold=True, color=NAVY); ini[f"B{i}"] = b; ini[f"B{i}"].alignment = wrap; ini.row_dimensions[i].height = 34

# ---- Campañas
cp = hoja("Campanas", ["Código (SIG)", "Campaña", "Universidad", "Fecha examen", "Cierre inscripción", "Inicio", "Fin", "Público", "Oferta y precios", "Gancho", "Canales", "Ads S/ por día", "Días de ads", "Presupuesto ads (S/)", "Meta leads", "Meta inscritos"],
          [16, 26, 12, 12, 12, 11, 11, 34, 40, 40, 26, 9, 8, 12, 9, 9], tab=GOLD)
for i, c in enumerate(CAMPANAS, 2):
    cod, nom, uni, ex, cierre, ini_, fin, pub, oferta, gancho, canales, ads, dias, ml, mi = c
    of = " · ".join(f"{PRECIO[p]['producto']} ({PRECIO[p]['horas']} h) {soles(p)}" for p in oferta) or "—"
    for j, v in enumerate([cod, nom, uni, ex, cierre, ini_, fin, pub, of, gancho, canales, ads, dias], 1):
        cp.cell(i, j, v); cp.cell(i, j).alignment = wrap
        if j in (4, 5, 6, 7) and v: cp.cell(i, j).number_format = "dd/mm/yyyy"
    cp.cell(i, 14, f"=L{i}*M{i}").number_format = '"S/ "#,##0'
    cp.cell(i, 15, ml); cp.cell(i, 16, mi)
    cp.row_dimensions[i].height = 60

# ---- Calendario semanal
cal = hoja("Calendario", ["Semana (lunes)", "Día", "Fecha", "Campaña", "Formato", "Fase", "Idea / texto sugerido", "Llamado a la acción", "Estado", "Notas"],
           [12, 10, 11, 16, 16, 14, 70, 34, 12, 20])
def fase(c, f):
    ex = c[3]
    if ex is None: return "Siempre"
    d = (ex - f).days
    return "Lanzamiento" if d > 21 else "Valor" if d > 7 else "Urgencia"
UNI_EX = {"ULIMA": "el Examen de Escolares ULIMA", "PUCP": "La Primera Opción PUCP", "UPCH": "el examen de Cayetano"}
def texto(c, f, formato):
    cod, nom, uni, ex = c[:4]
    fz = fase(c, f); dias = (ex - f).days if ex else None
    examen = {"SPRINT-ULIMA": "el Examen de Escolares ULIMA", "SPRINT-PUCP": "La Primera Opción PUCP", "REFUERZO-CEPU": "el 3.er examen del CEPU",
              "MEDICINA-UPCH": "Medicina en Cayetano", "VERANO-PUCP": "la Evaluación del Talento PUCP", "VERANO-ULIMA": "el examen de egresados ULIMA",
              "CACHIMBOS": "tu primer ciclo"}.get(cod, "")
    if cod == "CACHIMBOS":
        return ("¿Ingresaste? Cálculo, Física y Química del primer ciclo no esperan. Únete a la lista de espera y asegura tu cupo con tutor."
                if f < D(2027, 2, 1) else "Primer ciclo sin sustos: practica tus EP y exámenes con un tutor que conoce el curso de tu universidad."), "Escribe 'CACHIMBO' al 902 155 458"
    k = f.isocalendar()[1]
    V = {
     ("Lanzamiento", "Reel"): [f"Reel: '¿Postulas a {examen}? En 45 minutos te digo qué temas te faltan.' Mostrar cómo es la clase diagnóstica y el plan por semanas.",
                               f"Reel: 'Un día de estudio con ASTROS': clase virtual, pizarra digital y grabación en Drive, rumbo a {examen}.",
                               f"Reel: 'Personalizado de verdad': tú eliges el temario y el horario para {examen}."],
     ("Lanzamiento", "Carrusel"): [f"Carrusel: 'Cómo es {examen}': áreas que evalúa y cómo prepararte en {max(dias // 7, 1)} semanas.",
                                   f"Carrusel: plan semana a semana para {examen} (qué estudiar cada semana).",
                                   f"Carrusel: 'Academia grande vs. tutoría personalizada' para {examen}: horarios, atención y avance medible."],
     ("Lanzamiento", "Historia"): [f"Historia con encuesta: '¿Ya sabes en qué área estás más débil para {examen}?' Sí / No → link a WhatsApp.",
                                   "Historia: caja de preguntas '¿Qué tema te da más miedo?' y responder con mini videos."],
     ("Lanzamiento", "WhatsApp"): [f"Difusión: guion 'Primer contacto' adaptado a {examen}. Ofrece el diagnóstico gratis.",
                                   "Difusión: guion 'Seguimiento' a quienes preguntaron y no agendaron diagnóstico."],
     ("Valor", "Reel"): [f"Reel: resolver en 60 segundos una pregunta tipo {uni} del banco ASTROS. Cerrar con 'quedan {dias} días'.",
                         f"Reel: 'El truco para no perder tiempo en Razonamiento Matemático' (pregunta del banco tipo {uni})."],
     ("Valor", "Carrusel"): [f"Carrusel: '3 errores que cuestan puntos en {examen}' (tiempo, lectura apurada, no revisar claves).",
                             "Carrusel: cómo leer un texto de comprensión lectora en 3 pasos (con una pregunta del banco)."],
     ("Valor", "Historia"): ["Historia: reto diario de 1 pregunta del banco con respuesta al día siguiente.",
                             "Historia: resultado del mini simulacro de la semana (promedio del grupo, sin nombres)."],
     ("Valor", "WhatsApp"): ["Difusión: guion 'Seguimiento' a leads que no se inscribieron + envío de mini simulacro PDF.",
                             "Difusión: invitación al simulacro gratuito del sábado (cupos limitados)."],
     ("Urgencia", "Reel"): [f"Reel: 'Faltan {dias} días para {examen}: esto es lo que haría yo esta semana.'"],
     ("Urgencia", "Carrusel"): ["Carrusel: checklist del día del examen (DNI, hora, descanso, estrategia de tiempo)."],
     ("Urgencia", "Historia"): [f"Historia: cuenta regresiva a {examen} + 'últimos cupos para el simulacro final'."],
     ("Urgencia", "WhatsApp"): ["Difusión: guion 'Simulacro final' + recordatorio de inscripción."],
    }[(fz, formato)]
    t = re.sub(r"\b([aA]) el ", lambda m: m.group(1) + "l ", V[k % len(V)])
    return t, "Escribe 'DIAGNÓSTICO' al 902 155 458"
ORG = ["Frase ASTROS de la semana + tip de estudio", "Testimonio de alumno (con permiso)", "Detrás de cámaras: cómo preparo una clase", "Mito o realidad sobre los exámenes de admisión"]
fila = 2
lunes = D(2026, 10, 12)
while lunes <= D(2027, 4, 26):
    activas = [c for c in CAMPANAS if c[3] and c[0] not in ("REFERIDOS", "ORGANICO") and c[5] <= lunes + dt.timedelta(days=6) and c[6] >= lunes]
    activas.sort(key=lambda c: c[3])
    plan = []
    if activas:
        a1 = activas[0]; a2 = activas[1] if len(activas) > 1 else activas[0]
        plan = [(0, a1, "Reel"), (2, a2, "Carrusel"), (4, a1, "Historia"), (5, a1, "WhatsApp")]
    else:
        plan = [(0, None, "Reel"), (2, None, "Carrusel"), (5, None, "WhatsApp")]
    for off, c, fmt in plan:
        f = lunes + dt.timedelta(days=off)
        if f < D(2026, 10, 13): continue
        if c and c[3] and f > c[3]: c = None
        if c is None:
            ref = lunes >= D(2026, 11, 30)
            cod = "REFERIDOS" if (fmt == "WhatsApp" and ref) else "ORGANICO"
            t = ("Difusión: guion 'Referidos' a alumnos y ex alumnos." if cod == "REFERIDOS" else ORG[(fila // 3) % len(ORG)])
            row = [lunes, ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"][off], f, cod, fmt, "Siempre", t, "Seguir a @1.masesores", "Pendiente", ""]
        else:
            t, cta = texto(c, f, fmt)
            row = [lunes, ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"][off], f, c[0], fmt, fase(c, f), t, cta, "Pendiente", ""]
        for j, v in enumerate(row, 1):
            cal.cell(fila, j, v); cal.cell(fila, j).alignment = wrap
        for j in (1, 3): cal.cell(fila, j).number_format = "dd/mm/yyyy"
        fila += 1
    lunes += dt.timedelta(days=7)
# hitos fijos
HITOS = [(D(2026, 10, 14), "SPRINT-ULIMA", "Historia", "Urgencia", "Rendimiento Académico ULIMA cierra el 16/10: 'Te preparo para la entrevista'.", "Escribe 'ULIMA' al 902 155 458"),
         (D(2026, 11, 3), "SPRINT-ULIMA", "WhatsApp", "Urgencia", "Recordatorio: la inscripción al examen ULIMA cierra el 06/11 a las 18:00.", "Responder 'LISTO' si ya te inscribiste"),
         (D(2026, 11, 30), "REFERIDOS", "Carrusel", "Siempre", "Publicar ingresantes ULIMA/PUCP (con permiso) + '1 hora gratis por cada amigo que traigas'.", "Escribe 'TRIPULACIÓN'"),
         (D(2027, 1, 7), "MEDICINA-UPCH", "WhatsApp", "Urgencia", "Recordatorio: inscripción a Medicina UPCH cierra el 11/01.", "Responder 'LISTO' si ya te inscribiste"),
         (D(2026, 11, 2), "ORGANICO", "Tarea", "Siempre", "Verificar fechas no publicadas (PUCP Talento, ULIMA egresados, UPC) y actualizar este plan.", "—"),
         (D(2027, 1, 5), "ORGANICO", "Tarea", "Siempre", "Verificar de nuevo fechas de PUCP Talento y ULIMA egresados.", "—")]
for f, cod, fmt, fz, t, cta in HITOS:
    lunes = f - dt.timedelta(days=f.weekday())
    row = [lunes, ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"][f.weekday()], f, cod, fmt, fz, "HITO: " + t, cta, "Pendiente", ""]
    for j, v in enumerate(row, 1):
        cal.cell(fila, j, v); cal.cell(fila, j).alignment = wrap
    for j in (1, 3): cal.cell(fila, j).number_format = "dd/mm/yyyy"
    fila += 1
# ordenar por fecha
rows = sorted([[cal.cell(r, j).value for j in range(1, 11)] for r in range(2, fila)], key=lambda r: (r[2], r[4] != "Tarea"))
for r, vals in enumerate(rows, 2):
    for j, v in enumerate(vals, 1): cal.cell(r, j, v)
cal.auto_filter.ref = f"A1:J{fila-1}"
dv = DataValidation(type="list", formula1='"Pendiente,Programado,Publicado,No se hizo"', allow_blank=True); cal.add_data_validation(dv); dv.add(f"I2:I{fila+200}")

# ---- WhatsApp
wa = hoja("WhatsApp", ["Guion", "Cuándo usarlo", "Texto (copia y pega, cambia lo que está entre [ ])"], [22, 30, 100], tab=GOLD)
GUIONES = [
 ("Primer contacto", "Alguien escribe por primera vez", "¡Hola [nombre]! 👋 Soy Martini de ASTROS Maestros Asesores. Gracias por escribirnos. ¿A qué universidad postulas y cuándo es tu examen? Te ofrezco una clase diagnóstica GRATIS de 45 minutos: vemos qué temas dominas, cuáles te faltan y armamos tu plan semana a semana. ¿Qué día te acomoda?"),
 ("Después del diagnóstico", "El mismo día del diagnóstico", "¡Gracias por la clase de hoy, [nombre]! 🚀 Te resumo: estás fuerte en [tema] y conviene reforzar [tema 1] y [tema 2]. Para llegar listo(a) a [examen] el [fecha], te recomiendo [producto] ([horas] h, [precio]). Puedes pagar por Yape, Plin o transferencia. ¿Te reservo tus horarios?"),
 ("Seguimiento (2-3 días)", "No respondió a la propuesta", "Hola [nombre], ¿pudiste revisar el plan? Te dejo un mini simulacro tipo [universidad] con solucionario para que practiques 📄. Si te animas, esta semana aún tengo horarios [días]."),
 ("Cierre", "Aceptó la propuesta", "¡Excelente, [nombre]! 🎉 Para confirmar tu cupo, el pago es de [monto] por Yape/Plin o transferencia. Envíame el comprobante por aquí y te mando tu horario y el acceso a tu carpeta de clases grabadas."),
 ("Recordatorio de inscripción", "Una semana antes del cierre oficial", "⏰ [nombre], la inscripción para [examen] cierra el [fecha y hora]. ¿Ya te inscribiste? Si necesitas ayuda con el proceso, avísame."),
 ("Simulacro final", "1-2 semanas antes del examen", "Hola [nombre], este [día] hacemos el simulacro final tipo [universidad], cronometrado y con corrección personalizada. ¿Te separo un cupo?"),
 ("Post-resultado: ingresó", "Día de resultados", "¡FELICITACIONES, [nombre]! 🎓🚀 Lo lograste. ¿Me das permiso para compartir tu logro en nuestras redes? Y si tienes amigos que postulan, por cada uno que se inscriba te regalo 1 hora de tutoría para tu primer ciclo."),
 ("Post-resultado: no ingresó", "Día de resultados", "[nombre], sé que no es la noticia que esperabas. Lo que preparaste no se pierde: el próximo proceso es [fecha]. Te propongo revisar juntos tu examen y armar un plan de verano. ¿Conversamos esta semana?"),
 ("Referidos", "Alumnos actuales y ex alumnos", "Hola [nombre] 👋 ¿Conoces a alguien que postule a PUCP, ULIMA, UPC o Cayetano, o que necesite ayuda en la universidad? Por cada amigo que se inscriba en un paquete, te regalo 1 hora de clase. ¡Gracias por ser parte de la tripulación ASTROS!"),
]
for i, g in enumerate(GUIONES, 2):
    for j, v in enumerate(g, 1): wa.cell(i, j, v); wa.cell(i, j).alignment = wrap
    wa.row_dimensions[i].height = 70

# ---- Presupuesto
ps = hoja("Presupuesto", ["Código", "Presupuesto ads (S/)", "Meta inscritos", "Costo meta por inscrito (S/)", "Gasto real (S/)", "Inscritos reales (del SIG)", "Costo real por inscrito (S/)"], [16, 14, 12, 16, 14, 16, 16], tab=RED)
for i in range(2, len(CAMPANAS) + 2):
    ps.cell(i, 1, f"=Campanas!A{i}"); ps.cell(i, 2, f"=Campanas!N{i}"); ps.cell(i, 3, f"=Campanas!P{i}")
    ps.cell(i, 4, f'=IFERROR(B{i}/C{i},"")'); ps.cell(i, 7, f'=IFERROR(E{i}/F{i},"")')
    for j in (2, 4, 5, 7): ps.cell(i, j).number_format = '"S/ "#,##0'
t = len(CAMPANAS) + 2
ps.cell(t, 1, "TOTAL").font = Font(bold=True)
for col in "BCEF": ps[f"{col}{t}"] = f"=SUM({col}2:{col}{t-1})"
ps[f"D{t}"] = f'=IFERROR(B{t}/C{t},"")'; ps[f"G{t}"] = f'=IFERROR(E{t}/F{t},"")'
for col in "BDEG": ps[f"{col}{t}"].number_format = '"S/ "#,##0'
ps.cell(t + 2, 1, "Anuncios en Meta (IG/FB): segmenta Lima, 16-19 años y padres 38-55, intereses 'admisión universitaria' y la universidad. Empieza con el mínimo diario y sube solo en los anuncios con más mensajes por sol.").alignment = wrap
ps.merge_cells(start_row=t + 2, start_column=1, end_row=t + 2, end_column=7); ps.row_dimensions[t + 2].height = 45

wb._sheets = [wb[n] for n in ["Inicio", "Campanas", "Calendario", "WhatsApp", "Presupuesto"]]
wb.save(OUT)
print(f"OK {OUT} · {fila-2} piezas en calendario")

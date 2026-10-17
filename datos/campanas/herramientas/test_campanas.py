"""Pruebas del plan de campañas: códigos iguales al SIG, precios existentes en el tarifario, fechas iguales
al cronograma de admisión, ninguna pieza de una campaña después de su examen y presupuesto bien sumado.
Uso: python3 test_campanas.py"""
import csv, datetime as dt, os, re, subprocess, sys, tempfile
from openpyxl import load_workbook
AQUI = os.path.dirname(os.path.abspath(__file__)); BASE = os.path.dirname(AQUI); RAIZ = os.path.dirname(BASE)
fallos = []
def check(n, ok, det=""):
    print("OK   " if ok else "FALLA", n, "" if ok else det)
    if not ok: fallos.append(n)
tmp = tempfile.mkdtemp(); out = os.path.join(tmp, "c.xlsx")
subprocess.run([sys.executable, os.path.join(AQUI, "construir_campanas.py"), out], check=True, capture_output=True)
wb = load_workbook(out)
camp = {r[0]: r for r in wb["Campanas"].iter_rows(min_row=2, values_only=True) if r[0]}
cal = [r for r in wb["Calendario"].iter_rows(min_row=2, values_only=True) if r[0]]

sig_src = open(os.path.join(RAIZ, "sig", "herramientas", "construir_sig.py"), encoding="utf-8").read()
cod_sig = set(re.search(r'"Campaña": \[(.*?)\]', sig_src).group(1).replace('"', "").replace(" ", "").split(","))
check("códigos de campaña iguales a los del SIG", set(camp) == cod_sig, f"{set(camp) ^ cod_sig}")
check("calendario solo usa códigos existentes", {r[3] for r in cal} <= set(camp))

tar = {r["id"]: r for r in csv.DictReader(open(os.path.join(RAIZ, "precios", "tarifario.csv"), encoding="utf-8")) if r["activo"] == "si"}
src = open(os.path.join(AQUI, "construir_campanas.py"), encoding="utf-8").read()
ids = set(re.findall(r'"((?:esc|pre|uni)-[a-z0-9]+)"', src))
check("productos ofrecidos existen y están activos en el tarifario", ids <= set(tar), f"{ids - set(tar)}")
for cod, r in camp.items():
    for m in re.finditer(r"S/ ([\d,]+)", r[8] or ""):
        check(f"{cod}: precio {m.group(0)} existe en el tarifario", any(float(t["precio_virtual"]) == float(m.group(1).replace(",", "")) for t in tar.values()))

adm = load_workbook(os.path.join(RAIZ, "admisiones", "ASTROS_calendario_admision_2027.xlsx"))["Fechas_admision"]
oficial = {}
for r in adm.iter_rows(min_row=2, values_only=True):
    if isinstance(r[6], dt.datetime) and r[8] == "Próximo": oficial.setdefault((r[0], r[3]), r[6].date())
def fecha(u, sub): return next(v for (uu, m), v in oficial.items() if uu == u and sub in m)
d = lambda x: x.date() if isinstance(x, dt.datetime) else x
check("Sprint ULIMA = fecha oficial del Examen de Escolares", d(camp["SPRINT-ULIMA"][3]) == fecha("ULIMA", "Examen de Escolares"))
check("Sprint PUCP = fecha oficial de La Primera Opción", d(camp["SPRINT-PUCP"][3]) == fecha("PUCP", "La Primera Opción"))
check("Refuerzo CEPU = fecha oficial CEPU II", d(camp["REFUERZO-CEPU"][3]) == fecha("UPCH", "CEPU II"))
check("Medicina UPCH = fecha oficial", d(camp["MEDICINA-UPCH"][3]) == fecha("UPCH", "MEDICINA"))
for cod, r in camp.items():
    if r[3] and r[4]: check(f"{cod}: cierre de inscripción antes del examen", d(r[4]) < d(r[3]))
    if r[5]: check(f"{cod}: inicio antes del fin", d(r[5]) < d(r[6]))

tardias = [(r[3], d(r[2])) for r in cal if camp[r[3]][3] and r[3] not in ("REFERIDOS", "ORGANICO") and d(r[2]) > d(camp[r[3]][3])]
check("ninguna pieza de campaña después de su examen", not tardias, str(tardias[:5]))
por_semana = {}
for r in cal: por_semana[d(r[0])] = por_semana.get(d(r[0]), 0) + 1
check("todas las semanas tienen al menos 3 piezas", min(por_semana.values()) >= 3)
semanas = sorted(por_semana); check("sin semanas vacías entre oct y abr", all((b - a).days == 7 for a, b in zip(semanas, semanas[1:])))
check("fechas en orden", [d(r[2]) for r in cal] == sorted(d(r[2]) for r in cal))
for txt in ("06/11", "11/01", "16/10"):
    check(f"hito de cierre {txt} presente", any(txt in (r[6] or "") for r in cal))
check("sin errores de contracción ('a el', 'paral')", not any(re.search(r"\ba el |paral ", r[6] or "") for r in cal))
check("ningún texto con datos sin reemplazar ({...})", not any("{" in (r[6] or "") for r in cal))

subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", os.path.join(tmp, "r"), out], check=True, capture_output=True)
rc = load_workbook(os.path.join(tmp, "r", "c.xlsx"), data_only=True)
ps = rc["Presupuesto"]; n = len(camp) + 2
esperado = sum((r[11] or 0) * (r[12] or 0) for r in camp.values())
check(f"presupuesto total = S/ {esperado}", ps[f"B{n}"].value == esperado, str(ps[f"B{n}"].value))
check("sin errores de fórmula", not any(isinstance(c.value, str) and c.value.startswith("#") for ws in rc for row in ws.iter_rows() for c in row))
print(f"\n{'TODAS LAS PRUEBAS PASARON' if not fallos else str(len(fallos)) + ' FALLA(S): ' + '; '.join(fallos)}")
sys.exit(1 if fallos else 0)

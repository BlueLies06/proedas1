"""Actualiza solo la hoja Catalogo del SIG desde precios/tarifario.csv y precios/reglas.csv.
Antes de tocar nada guarda un respaldo en sig/backups/ y anota los cambios en sig/logs/cambios.log.
Uso: python3 sincronizar_catalogo.py [ruta/ASTROS_SIG.xlsx]"""
import csv, datetime as dt, os, shutil, sys
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PRECIOS = os.environ.get("PRECIOS_DIR", os.path.join(os.path.dirname(BASE), "precios"))
SIG = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "ASTROS_SIG.xlsx")
NIV = {"escolar": "Escolar", "pre": "Preuniversitario", "universitario": "Universitario"}

def leer_catalogo():
    reglas = {r["regla"]: float(r["valor"]) for r in csv.DictReader(open(os.path.join(PRECIOS, "reglas.csv"), encoding="utf-8"))}
    filas = []
    for r in csv.DictReader(open(os.path.join(PRECIOS, "tarifario.csv"), encoding="utf-8")):
        if r["activo"].strip().lower() != "si":
            continue
        n = int(r["alumnos_max"])
        tutor = (float(r["pago_tutor_hora"]) + reglas["extra_tutor_por_alumno"] * (n - 1)) / n
        filas.append((r["id"], NIV.get(r["nivel"], r["nivel"]), r["producto"], float(r["horas"]), float(r["precio_virtual"]), round(tutor, 2)))
    filas += [("ext-presencial", "Todos", "Recargo presencial, por hora", 1, reglas["recargo_presencial_hora"], 0),
              ("ext-domicilio", "Todos", "Recargo a domicilio, por hora", 1, reglas["recargo_domicilio_hora"], 0),
              ("diagnostico", "Todos", "Clase diagnóstica gratis", reglas["clase_diagnostica_min"] / 60, 0, 0)]
    return filas

def main():
    ahora = dt.datetime.now()
    wb = load_workbook(SIG)
    ws = wb["Catalogo"]
    antes = {ws.cell(i, 1).value: (ws.cell(i, 5).value, ws.cell(i, 4).value, ws.cell(i, 7).value)
             for i in range(2, 61) if ws.cell(i, 1).value}
    nuevo = leer_catalogo()
    if len(nuevo) > 59:
        sys.exit("ERROR: el catálogo admite hasta 59 productos")
    despues = {c: (pr, h, t) for c, _, _, h, pr, t in nuevo}
    cambios = []
    for c in sorted(set(antes) | set(despues)):
        if c not in antes: cambios.append(f"{c}: nuevo {despues[c]}")
        elif c not in despues: cambios.append(f"{c}: retirado (antes {antes[c]})")
        elif tuple(float(x or 0) for x in antes[c]) != tuple(float(x) for x in despues[c]):
            cambios.append(f"{c}: (precio, horas, pago tutor) {antes[c]} -> {despues[c]}")
    os.makedirs(os.path.join(BASE, "logs"), exist_ok=True)
    log = open(os.path.join(BASE, "logs", "cambios.log"), "a", encoding="utf-8")
    if not cambios:
        log.write(f"{ahora:%Y-%m-%d %H:%M:%S} sincronizar_catalogo: sin cambios\n"); print("Sin cambios"); return
    usados = {wb["Ventas"].cell(i, 5).value for i in range(2, wb["Ventas"].max_row + 1)} - {None}
    retirados_usados = sorted(c for c in usados if c not in despues)
    destino = os.path.join(BASE, "backups", f"{ahora:%Y%m%d-%H%M%S}")
    os.makedirs(destino, exist_ok=True)
    shutil.copy2(SIG, destino)
    for i in range(2, 61):
        for j in (1, 2, 3, 4, 5, 7, 9): ws.cell(i, j).value = None
        ws.cell(i, 6).value = ws.cell(i, 8).value = None
    for i, (c, n, p, h, pr, t) in enumerate(nuevo, 2):
        for j, v in enumerate((c, n, p, h, pr), 1): ws.cell(i, j, v)
        ws.cell(i, 6, f'=IFERROR(E{i}/D{i},"")'); ws.cell(i, 7, t)
        ws.cell(i, 8, f'=IFERROR((F{i}-G{i})/F{i},"")'); ws.cell(i, 9, "precios/tarifario.csv")
    wb.save(SIG)
    log.write(f"{ahora:%Y-%m-%d %H:%M:%S} sincronizar_catalogo: respaldo {destino}\n")
    for c in cambios: log.write(f"    {c}\n")
    if retirados_usados:
        msg = f"AVISO: ventas usan códigos retirados del tarifario: {', '.join(retirados_usados)} (sus horas y precio quedarán vacíos)"
        log.write(f"    {msg}\n"); print(msg)
    print(f"{len(cambios)} cambio(s). Respaldo en {destino}")

if __name__ == "__main__":
    main()

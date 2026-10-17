"""Pruebas automáticas del SIG: construye una copia con datos de prueba (incluye errores a propósito),
la recalcula con LibreOffice y verifica cálculos, alertas de la hoja Control y la sincronización del catálogo.
Uso: python3 test_sig.py   (necesita LibreOffice: soffice)"""
import csv, datetime as dt, os, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
TMP = tempfile.mkdtemp(prefix="test_sig_")
fallos = []
def check(nombre, real, esperado):
    ok = real == esperado or (isinstance(esperado, float) and real is not None and abs(real - esperado) < 1e-6)
    print(("OK   " if ok else "FALLA"), nombre, "" if ok else f"(esperado {esperado!r}, obtuvo {real!r})")
    if not ok: fallos.append(nombre)

def recalcular(ruta):
    out = os.path.join(TMP, "recalc")
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, ruta], check=True, capture_output=True)
    return load_workbook(os.path.join(out, os.path.basename(ruta)), data_only=True)

base = os.path.join(TMP, "sig.xlsx")
subprocess.run([sys.executable, os.path.join(AQUI, "construir_sig.py"), base], check=True, capture_output=True)
wb = load_workbook(base)
hoy = dt.date.today()
L, A, V, P, C = (wb[n] for n in ("Leads", "Alumnos", "Ventas", "Pagos", "Clases"))
for r, fila in enumerate([(hoy, "Ana", "PUCP", "Instagram", "SPRINT-PUCP", "Inscrito", None),
                          (hoy, "Luis", "ULIMA", "WhatsApp", "SPRINT-ULIMA", "Contactado", hoy - dt.timedelta(days=3)),
                          (hoy, "Rosa", "UPCH", "Web", "MEDICINA-UPCH", "Inscrito", None)], 2):
    for c, v in zip("BCFHIJK", fila): L[f"{c}{r}"] = v
A["B2"], A["L2"], A["M2"] = "Ana", "L001", "Activo"
A["B3"], A["M3"] = "Beto", "Activo"
A["B4"], A["M4"] = "Carla", "Activo"
V["B2"], V["C2"], V["E2"], V["H2"], V["M2"] = hoy, "A001", "pre-pack12", 40, "SPRINT-PUCP"
V["B3"], V["C3"], V["E3"] = hoy - dt.timedelta(days=90), "A002", "uni-practica"
V["B4"], V["C4"], V["E4"] = hoy, "A009", "no-existe"
V["B5"], V["C5"], V["E5"] = hoy - dt.timedelta(days=70), "A003", "esc-pack8"
for r, fila in enumerate([(hoy, "OP1", "A001", "V001", 300, "Yape", "Sí"),
                          (hoy, "OP1", "A001", "V001", 100, "Yape", "No"),
                          (hoy, "OP3", "A002", None, 50, "Plin", "Sí")], 2):
    for c, v in zip("ABCDEFG", fila): P[f"{c}{r}"] = v
for r, fila in enumerate([(hoy, "A001", "Martini", 2, "Asistió"), (hoy, "A001", "Martini", 1.5, "Reprogramada"),
                          (hoy, "A002", "Martini", 6, "Asistió"), (hoy, "A002", "Martini", None, "Asistió")], 2):
    for c, v in zip("ACFIJ", fila): C[f"{c}{r}"] = v
prueba = os.path.join(TMP, "prueba.xlsx"); wb.save(prueba)
r = recalcular(prueba)

al, ve, db, ctl = r["Alumnos"], r["Ventas"], r["Dashboard"], r["Control"]
check("ID automático de alumno", al["A2"].value, "A001")
check("Horas compradas Ana", al["N2"].value, 12)
check("Reprogramada no descuenta horas", al["O2"].value, 2)
check("Horas restantes Ana", al["P2"].value, 10)
check("Total vendido Ana (540-40)", al["Q2"].value, 500)
check("Deuda Ana (500-400)", al["S2"].value, 100)
check("Horas negativas Beto (4-6)", al["P3"].value, -2)
check("Saldo vacío si el producto no existe", ve["K4"].value, None)
check("Precio desde catálogo", ve["G2"].value, 540)
check("Vencimiento a 60 días", ve["L2"].value.date(), hoy + dt.timedelta(days=60))
check("Dashboard ingresos del mes", db["C8"].value, 450)
check("Dashboard ventas del mes (excluye la de hace 70 días)", db["C10"].value, 2)
check("Dashboard leads nuevos", db["C14"].value, 3)
check("Dashboard conversión", db["C16"].value, 2 / 3)
esperado = {2: 1, 3: 1, 4: 0, 5: 0, 6: 1, 7: 2, 8: 1, 9: 0, 10: 1, 11: 0, 12: 1, 13: 1, 14: 1}
for fila, n in esperado.items():
    check(f"Control: {ctl.cell(fila, 1).value}", ctl.cell(fila, 2).value, n)
check("Dashboard total de alertas", db["F15"].value, sum(esperado.values()))
for ws in r:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper():
                fallos.append(f"error {c.value} en {ws.title}!{c.coordinate}")

# sincronización: copiar tarifario, cambiar un precio y retirar un producto
pdir = os.path.join(TMP, "precios"); os.makedirs(pdir)
real = os.environ.get("PRECIOS_DIR", os.path.join(os.path.dirname(os.path.dirname(AQUI)), "precios"))
shutil.copy(os.path.join(real, "reglas.csv"), pdir)
filas = list(csv.DictReader(open(os.path.join(real, "tarifario.csv"), encoding="utf-8")))
for f in filas:
    if f["id"] == "pre-pack12": f["precio_virtual"] = "600"
    if f["id"] == "uni-practica": f["activo"] = "no"
with open(os.path.join(pdir, "tarifario.csv"), "w", newline="", encoding="utf-8") as fh:
    w = csv.DictWriter(fh, fieldnames=filas[0].keys()); w.writeheader(); w.writerows(filas)
sigdir = os.path.join(TMP, "sig"); os.makedirs(os.path.join(sigdir, "herramientas"))
shutil.copy(prueba, os.path.join(sigdir, "ASTROS_SIG.xlsx"))
shutil.copy(os.path.join(AQUI, "sincronizar_catalogo.py"), os.path.join(sigdir, "herramientas"))
out = subprocess.run([sys.executable, os.path.join(sigdir, "herramientas", "sincronizar_catalogo.py")],
                     env={**os.environ, "PRECIOS_DIR": pdir}, capture_output=True, text=True)
check("Sincronización termina bien", out.returncode, 0)
check("Avisa código retirado en uso", "uni-practica" in out.stdout, True)
check("Crea respaldo", len(os.listdir(os.path.join(sigdir, "backups"))), 1)
log = open(os.path.join(sigdir, "logs", "cambios.log"), encoding="utf-8").read()
check("Log registra cambio de precio", "pre-pack12" in log and "600" in log, True)
r2 = recalcular(os.path.join(sigdir, "ASTROS_SIG.xlsx"))
check("Nuevo precio llega a Ventas", r2["Ventas"]["G2"].value, 600)
check("Datos del alumno se conservan", r2["Alumnos"]["B2"].value, "Ana")
out2 = subprocess.run([sys.executable, os.path.join(sigdir, "herramientas", "sincronizar_catalogo.py")],
                      env={**os.environ, "PRECIOS_DIR": pdir}, capture_output=True, text=True)
check("Segunda sincronización sin cambios no respalda", len(os.listdir(os.path.join(sigdir, "backups"))), 1)

shutil.rmtree(TMP, ignore_errors=True)
print(f"\n{'TODAS LAS PRUEBAS PASARON' if not fallos else str(len(fallos)) + ' FALLA(S): ' + '; '.join(fallos)}")
sys.exit(1 if fallos else 0)

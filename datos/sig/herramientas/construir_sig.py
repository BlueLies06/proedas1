import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
SAMPLE = len(sys.argv) > 2 and sys.argv[2] == "sample"
NAVY, GOLD, TEAL, RED, WHITE, LIGHT = "0B2533", "F9C74F", "1FA2A6", "E63946", "FFFFFF", "EEF6F7"
N = 500  # filas preparadas por hoja

wb = Workbook()
hdr_font = Font(bold=True, color=NAVY)
hdr_fill = PatternFill("solid", fgColor=GOLD)
calc_fill = PatternFill("solid", fgColor=LIGHT)
title_font = Font(bold=True, size=16, color=WHITE)
navy_fill = PatternFill("solid", fgColor=NAVY)
thin = Side(style="thin", color="CCCCCC")

def sheet(name, cols, widths, tab=TEAL, calc=()):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(1, i, c)
        cell.font, cell.fill = hdr_font, hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in calc:
        ws.cell(1, i).fill = PatternFill("solid", fgColor=TEAL)
        ws.cell(1, i).font = Font(bold=True, color=WHITE)
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32
    return ws

def dv_list(ws, rng, src):
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    ws.add_data_validation(dv); dv.add(rng)

# ---------- Listas ----------
L = {
 "Nivel": ["Escolar", "Preuniversitario", "Universitario", "Todos"],
 "Universidad": ["PUCP", "ULIMA", "UPC", "UPCH", "Otra", "No aplica"],
 "Canal": ["Instagram", "Facebook", "TikTok", "WhatsApp", "Web", "Referido", "Colegio", "Otro"],
 "EstadoLead": ["Nuevo", "Contactado", "Diagnóstico agendado", "Diagnóstico hecho", "Propuesta enviada", "Inscrito", "Perdido"],
 "Modalidad": ["Virtual", "Presencial", "A domicilio"],
 "MedioPago": ["Yape", "Plin", "BCP", "BBVA", "Interbank", "Scotiabank", "Efectivo"],
 "EstadoAlumno": ["Activo", "Pausado", "Terminó", "Ingresó"],
 "Asistencia": ["Asistió", "Faltó (se descuenta)", "Reprogramada"],
 "Campaña": ["SPRINT-ULIMA", "SPRINT-PUCP", "REFUERZO-CEPU", "MEDICINA-UPCH", "VERANO-PUCP", "VERANO-ULIMA", "CACHIMBOS", "REFERIDOS", "ORGANICO"],
}
ls = wb.create_sheet("Listas"); ls.sheet_properties.tabColor = "999999"
refs = {}
for j, (k, vals) in enumerate(L.items(), 1):
    col = get_column_letter(j)
    ls.cell(1, j, k).font = hdr_font; ls.cell(1, j).fill = hdr_fill
    for i, v in enumerate(vals, 2): ls.cell(i, j, v)
    ls.column_dimensions[col].width = 22
    refs[k] = f"Listas!${col}$2:${col}${len(vals)+1}"

# ---------- Catálogo ----------
cat = sheet("Catalogo", ["Código", "Nivel", "Producto", "Horas", "Precio (S/)", "S/ por hora", "Pago tutor por hora y alumno (S/)", "Margen % (grupo lleno)", "Origen (no editar aquí)"],
            [10, 16, 46, 8, 12, 12, 14, 10, 22], tab=GOLD, calc=(6, 8))
import csv, os
PRECIOS = os.environ.get("PRECIOS_DIR", os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "precios"))
reglas = {r["regla"]: float(r["valor"]) for r in csv.DictReader(open(f"{PRECIOS}/reglas.csv", encoding="utf-8"))}
NIV = {"escolar": "Escolar", "pre": "Preuniversitario", "universitario": "Universitario"}
catalog = []
for r in csv.DictReader(open(f"{PRECIOS}/tarifario.csv", encoding="utf-8")):
    if r["activo"].strip().lower() != "si": continue
    n = int(r["alumnos_max"])
    tutor = (float(r["pago_tutor_hora"]) + reglas["extra_tutor_por_alumno"] * (n - 1)) / n
    catalog.append((r["id"], NIV.get(r["nivel"], r["nivel"]), r["producto"], float(r["horas"]), float(r["precio_virtual"]), round(tutor, 2)))
catalog += [("ext-presencial", "Todos", "Recargo presencial, por hora", 1, reglas["recargo_presencial_hora"], 0),
            ("ext-domicilio", "Todos", "Recargo a domicilio, por hora", 1, reglas["recargo_domicilio_hora"], 0),
            ("diagnostico", "Todos", "Clase diagnóstica gratis", reglas["clase_diagnostica_min"] / 60, 0, 0)]
for i, (c, n, p, h, pr, t) in enumerate(catalog, 2):
    for j, v in enumerate((c, n, p, h, pr), 1): cat.cell(i, j, v)
    cat.cell(i, 6, f'=IFERROR(E{i}/D{i},"")'); cat.cell(i, 7, t)
    cat.cell(i, 8, f'=IFERROR((F{i}-G{i})/F{i},"")'); cat.cell(i, 9, "precios/tarifario.csv")
    cat.cell(i, 6).number_format = '"S/ "0.00'; cat.cell(i, 8).number_format = "0%"
    for k in (6, 8): cat.cell(i, k).fill = calc_fill
for i in range(2, 61): cat.cell(i, 5).number_format = '"S/ "#,##0'
CATR = "Catalogo!$A$2:$H$60"

# ---------- Leads ----------
ld = sheet("Leads", ["ID", "Fecha", "Nombre", "WhatsApp", "Nivel", "Universidad objetivo", "Carrera / curso", "Canal", "Código campaña", "Estado", "Próximo seguimiento", "Notas", "Días sin seguimiento"],
           [8, 12, 24, 14, 16, 14, 20, 12, 16, 20, 14, 30, 12], calc=(1, 13))
for r in range(2, N+2):
    ld.cell(r, 1, f'=IF(C{r}="","","L"&TEXT(ROW()-1,"000"))')
    ld.cell(r, 13, f'=IF(OR(K{r}="",J{r}="Inscrito",J{r}="Perdido"),"",TODAY()-K{r})')
    for c in (2, 11): ld.cell(r, c).number_format = "dd/mm/yyyy"
    for c in (1, 13): ld.cell(r, c).fill = calc_fill
for col, k in (("E", "Nivel"), ("F", "Universidad"), ("H", "Canal"), ("I", "Campaña"), ("J", "EstadoLead")):
    dv_list(ld, f"{col}2:{col}{N+1}", refs[k])
ld.conditional_formatting.add(f"M2:M{N+1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F8C9CC")))

# ---------- Alumnos ----------
al = sheet("Alumnos", ["ID", "Nombre", "WhatsApp", "Correo", "Nivel", "Colegio / Universidad actual", "Universidad objetivo", "Carrera", "Fecha examen", "Apoderado", "WhatsApp apoderado", "ID lead", "Estado", "Horas compradas", "Horas dictadas", "Horas restantes", "Total vendido (S/)", "Total pagado (S/)", "Deuda (S/)"],
           [8, 24, 14, 22, 16, 22, 14, 18, 12, 20, 14, 8, 10, 10, 10, 10, 12, 12, 12], calc=(1, 14, 15, 16, 17, 18, 19))
for r in range(2, N+2):
    al.cell(r, 1, f'=IF(B{r}="","","A"&TEXT(ROW()-1,"000"))')
    al.cell(r, 14, f'=IF(A{r}="","",SUMIFS(Ventas!$F:$F,Ventas!$C:$C,A{r}))')
    al.cell(r, 15, f'=IF(A{r}="","",SUMIFS(Clases!$I:$I,Clases!$C:$C,A{r},Clases!$J:$J,"<>Reprogramada"))')
    al.cell(r, 16, f'=IF(A{r}="","",N{r}-O{r})')
    al.cell(r, 17, f'=IF(A{r}="","",SUMIFS(Ventas!$I:$I,Ventas!$C:$C,A{r}))')
    al.cell(r, 18, f'=IF(A{r}="","",SUMIFS(Pagos!$E:$E,Pagos!$C:$C,A{r}))')
    al.cell(r, 19, f'=IF(A{r}="","",Q{r}-R{r})')
    al.cell(r, 9).number_format = "dd/mm/yyyy"
    for c in (17, 18, 19): al.cell(r, c).number_format = '"S/ "#,##0.00'
    for c in (1, 14, 15, 16, 17, 18, 19): al.cell(r, c).fill = calc_fill
for col, k in (("E", "Nivel"), ("G", "Universidad"), ("M", "EstadoAlumno")):
    dv_list(al, f"{col}2:{col}{N+1}", refs[k])
al.conditional_formatting.add(f"P2:P{N+1}", FormulaRule(formula=[f'AND(P2<>"",P2<=2)'], fill=PatternFill("solid", fgColor="FFE8A3")))
al.conditional_formatting.add(f"S2:S{N+1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F8C9CC")))

# ---------- Ventas ----------
ve = sheet("Ventas", ["ID venta", "Fecha", "ID alumno", "Alumno", "Código producto", "Horas", "Precio lista (S/)", "Descuento (S/)", "Total (S/)", "Pagado (S/)", "Saldo (S/)", "Vence (60 días)", "Código campaña", "Notas", "Vencido con horas pendientes"],
           [9, 12, 10, 22, 12, 8, 12, 12, 12, 12, 12, 12, 16, 28, 12], calc=(1, 4, 6, 7, 9, 10, 11, 12, 15))
for r in range(2, N+2):
    ve.cell(r, 1, f'=IF(C{r}="","","V"&TEXT(ROW()-1,"000"))')
    ve.cell(r, 4, f'=IFERROR(VLOOKUP(C{r},Alumnos!$A:$B,2,FALSE),"")')
    ve.cell(r, 6, f'=IFERROR(VLOOKUP(E{r},{CATR},4,FALSE),"")')
    ve.cell(r, 7, f'=IFERROR(VLOOKUP(E{r},{CATR},5,FALSE),"")')
    ve.cell(r, 9, f'=IF(G{r}="","",G{r}-H{r})')
    ve.cell(r, 10, f'=IF(A{r}="","",SUMIFS(Pagos!$E:$E,Pagos!$D:$D,A{r}))')
    ve.cell(r, 11, f'=IF(OR(A{r}="",I{r}=""),"",I{r}-J{r})')
    ve.cell(r, 15, f'=IF(L{r}="","",IF(AND(L{r}<TODAY(),IFERROR(VLOOKUP(C{r},Alumnos!$A:$P,16,FALSE),0)>0),"Sí","No"))')
    ve.cell(r, 12, f'=IF(B{r}="","",B{r}+60)')
    for c in (2, 12): ve.cell(r, c).number_format = "dd/mm/yyyy"
    for c in (7, 8, 9, 10, 11): ve.cell(r, c).number_format = '"S/ "#,##0.00'
    for c in (1, 4, 6, 7, 9, 10, 11, 12, 15): ve.cell(r, c).fill = calc_fill
dv_list(ve, f"E2:E{N+1}", "Catalogo!$A$2:$A$60")
dv_list(ve, f"M2:M{N+1}", refs["Campaña"])
ve.conditional_formatting.add(f"K2:K{N+1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F8C9CC")))

# ---------- Clases ----------
cl = sheet("Clases", ["Fecha", "Hora inicio", "ID alumno", "Alumno", "ID venta", "Tutor", "Curso", "Tema", "Horas", "Asistencia", "Modalidad", "Link grabación (Drive)", "Tarea / próximo paso", "Pago tutor (S/)"],
           [12, 10, 10, 22, 9, 16, 18, 26, 8, 18, 12, 30, 28, 12], calc=(4, 14))
for r in range(2, N+2):
    cl.cell(r, 4, f'=IFERROR(VLOOKUP(C{r},Alumnos!$A:$B,2,FALSE),"")')
    cl.cell(r, 14, f'=IF(OR(F{r}="",I{r}="",J{r}="Reprogramada"),"",I{r}*IFERROR(VLOOKUP(F{r},Tutores!$A:$F,6,FALSE),0))')
    cl.cell(r, 1).number_format = "dd/mm/yyyy"; cl.cell(r, 2).number_format = "hh:mm"
    cl.cell(r, 14).number_format = '"S/ "#,##0.00'
    for c in (4, 14): cl.cell(r, c).fill = calc_fill
dv_list(cl, f"J2:J{N+1}", refs["Asistencia"]); dv_list(cl, f"K2:K{N+1}", refs["Modalidad"])
dv_list(cl, f"F2:F{N+1}", "Tutores!$A$2:$A$50")

# ---------- Pagos ----------
pg = sheet("Pagos", ["Fecha", "N° operación", "ID alumno", "ID venta", "Monto (S/)", "Medio", "¿Comprobante recibido?", "Alumno", "Notas"],
           [12, 16, 10, 9, 12, 12, 12, 22, 28], calc=(8,))
for r in range(2, N+2):
    pg.cell(r, 8, f'=IFERROR(VLOOKUP(C{r},Alumnos!$A:$B,2,FALSE),"")'); pg.cell(r, 8).fill = calc_fill
    pg.cell(r, 1).number_format = "dd/mm/yyyy"; pg.cell(r, 5).number_format = '"S/ "#,##0.00'
dv_list(pg, f"F2:F{N+1}", refs["MedioPago"]); dv_list(pg, f"G2:G{N+1}", '"Sí,No"')

# ---------- Tutores ----------
tu = sheet("Tutores", ["Nombre", "WhatsApp", "Cursos que dicta", "Niveles", "Disponibilidad", "Pago por hora (S/)", "Estado", "Horas dictadas", "Por pagar (S/)"],
           [20, 14, 30, 18, 24, 12, 10, 10, 12], calc=(8, 9))
tu.cell(2, 1, "Martini"); tu.cell(2, 3, "Todos"); tu.cell(2, 4, "Todos"); tu.cell(2, 6, 0); tu.cell(2, 7, "Activo")
for r in range(2, 51):
    tu.cell(r, 8, f'=IF(A{r}="","",SUMIFS(Clases!$I:$I,Clases!$F:$F,A{r},Clases!$J:$J,"<>Reprogramada"))')
    tu.cell(r, 9, f'=IF(A{r}="","",SUMIFS(Clases!$N:$N,Clases!$F:$F,A{r}))')
    tu.cell(r, 9).number_format = '"S/ "#,##0.00'
    for c in (8, 9): tu.cell(r, c).fill = calc_fill

# ---------- Avance ----------
av = sheet("Avance", ["Fecha", "ID alumno", "Alumno", "Universidad", "Tipo (simulacro / práctica / nota univ.)", "Área / curso", "Puntaje", "Puntaje máximo", "% logro", "Temas débiles", "Comentario para el alumno/padre"],
           [12, 10, 22, 12, 20, 18, 10, 10, 9, 30, 34], calc=(3, 9))
for r in range(2, N+2):
    av.cell(r, 3, f'=IFERROR(VLOOKUP(B{r},Alumnos!$A:$B,2,FALSE),"")')
    av.cell(r, 9, f'=IF(OR(G{r}="",H{r}=""),"",G{r}/H{r})'); av.cell(r, 9).number_format = "0%"
    av.cell(r, 1).number_format = "dd/mm/yyyy"
    for c in (3, 9): av.cell(r, c).fill = calc_fill
dv_list(av, f"D2:D{N+1}", refs["Universidad"])

# ---------- Dashboard ----------
db = wb.create_sheet("Dashboard", 0); db.sheet_properties.tabColor = RED
db.column_dimensions["A"].width = 2
for c, w in zip("BCDEFG", (34, 16, 4, 30, 16, 16)): db.column_dimensions[c].width = w
db.merge_cells("B1:G1"); db["B1"] = "ASTROS · Panel de gestión"; db["B1"].font = title_font
for c in "BCDEFG": db[f"{c}1"].fill = navy_fill
db.row_dimensions[1].height = 34
db["B3"] = "Mes a analizar (cualquier día del mes):"; db["B3"].font = Font(bold=True)
db["C3"] = "=TODAY()"; db["C3"].number_format = "mmmm yyyy"; db["C3"].fill = hdr_fill
db["B4"] = "Inicio de mes"; db["C4"] = "=DATE(YEAR(C3),MONTH(C3),1)"
db["B5"] = "Fin de mes"; db["C5"] = "=EOMONTH(C3,0)"
for c in ("C4", "C5"): db[c].number_format = "dd/mm/yyyy"; db[c].font = Font(color="888888")
def kpi(row, col, label, formula, fmt=None):
    db[f"{col}{row}"] = label
    vc = chr(ord(col) + 1)
    db[f"{vc}{row}"] = formula; db[f"{vc}{row}"].font = Font(bold=True, size=12, color=NAVY)
    db[f"{vc}{row}"].fill = calc_fill
    if fmt: db[f"{vc}{row}"].number_format = fmt
S = '"S/ "#,##0'
for col, title in (("B", "Este mes"), ("E", "Acumulado")):
    db[f"{col}7"] = title; db[f"{col}7"].font = Font(bold=True, color=WHITE); db[f"{col}7"].fill = PatternFill("solid", fgColor=TEAL)
kpi(8, "B", "Ingresos cobrados", '=SUMIFS(Pagos!$E:$E,Pagos!$A:$A,">="&C4,Pagos!$A:$A,"<="&C5)', S)
kpi(9, "B", "Ventas cerradas (S/)", '=SUMIFS(Ventas!$I:$I,Ventas!$B:$B,">="&C4,Ventas!$B:$B,"<="&C5)', S)
kpi(10, "B", "N° de ventas", '=COUNTIFS(Ventas!$B:$B,">="&C4,Ventas!$B:$B,"<="&C5)')
kpi(11, "B", "Ticket promedio", '=IFERROR(C9/C10,0)', S)
kpi(12, "B", "Horas dictadas", '=SUMIFS(Clases!$I:$I,Clases!$A:$A,">="&C4,Clases!$A:$A,"<="&C5,Clases!$J:$J,"<>Reprogramada")', "0.0")
kpi(13, "B", "Pago a tutores", '=SUMIFS(Clases!$N:$N,Clases!$A:$A,">="&C4,Clases!$A:$A,"<="&C5)', S)
kpi(14, "B", "Leads nuevos", '=COUNTIFS(Leads!$B:$B,">="&C4,Leads!$B:$B,"<="&C5)')
kpi(15, "B", "Leads del mes ya inscritos", '=COUNTIFS(Leads!$B:$B,">="&C4,Leads!$B:$B,"<="&C5,Leads!$J:$J,"Inscrito")')
kpi(16, "B", "Conversión lead → alumno", '=IFERROR(C15/C14,0)', "0%")
kpi(8, "E", "Alumnos activos", '=COUNTIF(Alumnos!$M:$M,"Activo")')
kpi(9, "E", "Deuda por cobrar", '=SUMIF(Ventas!$K:$K,">0")', S)
kpi(10, "E", "Horas pendientes de dictar", '=SUMIF(Alumnos!$P:$P,">0")', "0.0")
kpi(11, "E", "Leads abiertos", '=COUNTA(Leads!$C:$C)-1-COUNTIF(Leads!$J:$J,"Inscrito")-COUNTIF(Leads!$J:$J,"Perdido")')
kpi(12, "E", "Seguimientos vencidos", '=COUNTIF(Leads!$M:$M,">0")')
kpi(13, "E", "Ingresantes", '=COUNTIF(Alumnos!$M:$M,"Ingresó")')
kpi(14, "E", "Ingresos totales cobrados", '=SUM(Pagos!$E:$E)', S)
kpi(15, "E", "Alertas en hoja Control", "=Control!B16")

r0 = 18
db[f"B{r0}"] = "Leads por canal (este mes)"; db[f"E{r0}"] = "Leads por universidad (este mes)"
for c in ("B", "E"):
    db[f"{c}{r0}"].font = Font(bold=True, color=WHITE); db[f"{c}{r0}"].fill = PatternFill("solid", fgColor=TEAL)
db[f"C{r0}"] = "Leads"; db[f"F{r0}"] = "Leads"; db[f"G{r0}"] = "Inscritos"
for i, ch in enumerate(L["Canal"], r0+1):
    db[f"B{i}"] = ch; db[f"C{i}"] = f'=COUNTIFS(Leads!$H:$H,B{i},Leads!$B:$B,">="&$C$4,Leads!$B:$B,"<="&$C$5)'
for i, u in enumerate(L["Universidad"][:5], r0+1):
    db[f"E{i}"] = u
    db[f"F{i}"] = f'=COUNTIFS(Leads!$F:$F,E{i},Leads!$B:$B,">="&$C$4,Leads!$B:$B,"<="&$C$5)'
    db[f"G{i}"] = f'=COUNTIFS(Leads!$F:$F,E{i},Leads!$B:$B,">="&$C$4,Leads!$B:$B,"<="&$C$5,Leads!$J:$J,"Inscrito")'
r1 = r0 + len(L["Canal"]) + 2
db[f"B{r1}"] = "Resultados por campaña (acumulado)"; db[f"B{r1}"].font = Font(bold=True, color=WHITE); db[f"B{r1}"].fill = PatternFill("solid", fgColor=TEAL)
db[f"C{r1}"] = "Leads"; db[f"E{r1}"] = "Inscritos"; db[f"F{r1}"] = "Ventas (S/)"; db[f"G{r1}"] = "Conversión"
for i, cp in enumerate(L["Campaña"], r1+1):
    db[f"B{i}"] = cp
    db[f"C{i}"] = f'=COUNTIF(Leads!$I:$I,B{i})'
    db[f"E{i}"] = f'=COUNTIFS(Leads!$I:$I,B{i},Leads!$J:$J,"Inscrito")'
    db[f"F{i}"] = f'=SUMIF(Ventas!$M:$M,B{i},Ventas!$I:$I)'; db[f"F{i}"].number_format = S
    db[f"G{i}"] = f'=IFERROR(E{i}/C{i},0)'; db[f"G{i}"].number_format = "0%"


# ---------- Cambios (registro) ----------
cb = sheet("Cambios", ["Fecha y hora", "Usuario", "Hoja", "Celda", "Valor anterior", "Valor nuevo", "Motivo (opcional)"],
           [18, 24, 12, 8, 24, 24, 30], tab="999999")
for r in range(2, 2001): cb.cell(r, 1).number_format = "dd/mm/yyyy hh:mm"

# ---------- Control de errores ----------
ce = wb.create_sheet("Control"); ce.sheet_properties.tabColor = RED
for c, w in zip("ABCD", (52, 10, 12, 60)): ce.column_dimensions[c].width = w
for j, h in enumerate(["Revisión automática", "Casos", "Estado", "Qué hacer"], 1):
    ce.cell(1, j, h).font = hdr_font; ce.cell(1, j).fill = hdr_fill
R = f"2:{N+1}"
checks = [
 ("Ventas con código de producto que no existe en el Catálogo", f"=SUMPRODUCT((Ventas!E2:E{N+1}<>\"\")*ISNA(MATCH(Ventas!E2:E{N+1},Catalogo!A2:A60,0)))", "Corrige el código o agrégalo al Catálogo."),
 ("Ventas con ID de alumno que no existe", f"=SUMPRODUCT((Ventas!C2:C{N+1}<>\"\")*ISNA(MATCH(Ventas!C2:C{N+1},Alumnos!A2:A{N+1},0)))", "Crea al alumno primero o corrige el ID."),
 ("Pagos con ID de venta que no existe", f"=SUMPRODUCT((Pagos!D2:D{N+1}<>\"\")*ISNA(MATCH(Pagos!D2:D{N+1},Ventas!A2:A{N+1},0)))", "Revisa el ID de venta del pago."),
 ("Pagos con ID de alumno que no existe", f"=SUMPRODUCT((Pagos!C2:C{N+1}<>\"\")*ISNA(MATCH(Pagos!C2:C{N+1},Alumnos!A2:A{N+1},0)))", "Revisa el ID de alumno del pago."),
 ("Pagos sin ID de venta (la deuda no baja)", f"=COUNTIFS(Pagos!E2:E{N+1},\">0\",Pagos!D2:D{N+1},\"\")", "Indica qué venta está pagando."),
 ("N° de operación repetido (posible pago duplicado)", f"=SUMPRODUCT((Pagos!B2:B{N+1}<>\"\")*(COUNTIF(Pagos!B2:B{N+1},Pagos!B2:B{N+1})>1))", "Borra el pago duplicado."),
 ("Pagos sin comprobante recibido", f"=COUNTIF(Pagos!G2:G{N+1},\"No\")", "Pide el comprobante al alumno."),
 ("Ventas pagadas de más (saldo negativo)", f"=COUNTIF(Ventas!K2:K{N+1},\"<0\")", "Revisa montos o asigna el exceso a otra venta."),
 ("Alumnos con horas negativas (dictaste más de lo vendido)", f"=COUNTIF(Alumnos!P2:P{N+1},\"<0\")", "Registra la venta de horas adicionales."),
 ("Clases con ID de alumno que no existe", f"=SUMPRODUCT((Clases!C2:C{N+1}<>\"\")*ISNA(MATCH(Clases!C2:C{N+1},Alumnos!A2:A{N+1},0)))", "Corrige el ID de alumno."),
 ("Clases sin horas registradas", f"=SUMPRODUCT((Clases!C2:C{N+1}<>\"\")*(Clases!I2:I{N+1}=\"\"))", "Anota cuántas horas duró la clase."),
 ("Paquetes vencidos con horas pendientes", f"=COUNTIF(Ventas!O2:O{N+1},\"Sí\")", "Contacta al alumno para usar o renovar sus horas."),
 ("Leads 'Inscrito' sin ficha de alumno (columna ID lead)", f"=SUMPRODUCT((Leads!J2:J{N+1}=\"Inscrito\")*(COUNTIF(Alumnos!L2:L{N+1},Leads!A2:A{N+1})=0))", "Crea la ficha del alumno y pon su ID de lead."),
]
for i, (t, f, a) in enumerate(checks, 2):
    ce.cell(i, 1, t); ce.cell(i, 2, f); ce.cell(i, 2).fill = calc_fill
    ce.cell(i, 3, f'=IF(B{i}=0,"OK","Revisar")'); ce.cell(i, 4, a)
    ce.cell(i, 4).alignment = Alignment(wrap_text=True)
last = len(checks) + 1
ce.conditional_formatting.add(f"C2:C{last}", CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="CDEFD8")))
ce.conditional_formatting.add(f"C2:C{last}", CellIsRule(operator="equal", formula=['"Revisar"'], fill=PatternFill("solid", fgColor="F8C9CC")))
ce.cell(last + 2, 1, "Total de alertas").font = Font(bold=True)
ce.cell(last + 2, 2, f"=SUM(B2:B{last})").font = Font(bold=True)

# ---------- Inicio ----------
ini = wb.create_sheet("Inicio", 0); ini.sheet_properties.tabColor = NAVY
ini.column_dimensions["A"].width = 2; ini.column_dimensions["B"].width = 22; ini.column_dimensions["C"].width = 90
ini.merge_cells("B1:C1"); ini["B1"] = "ASTROS · Sistema de Información Gerencial"; ini["B1"].font = title_font
ini["B1"].fill = navy_fill; ini["C1"].fill = navy_fill; ini.row_dimensions[1].height = 34
ini["B2"] = '"Aprender para crecer" · Maestros Asesores'; ini["B2"].font = Font(italic=True, color=TEAL)
rows = [
 ("Cómo funciona", "Escribe solo en las columnas con encabezado AMARILLO. Las de encabezado TURQUESA se calculan solas (no las borres)."),
 ("1. Leads", "Cada persona que escribe por IG, FB, WhatsApp o la web. Anota el canal y el código de campaña para saber qué funciona. Pon fecha de próximo seguimiento: si se vence, se pinta en rojo."),
 ("2. Alumnos", "Cuando un lead se inscribe, créalo aquí (el ID A001, A002... aparece solo al poner el nombre). Cambia el lead a 'Inscrito'. Aquí ves sus horas restantes y su deuda."),
 ("3. Ventas", "Cada paquete vendido: elige el ID del alumno y el código del producto (del Catálogo); horas y precio se llenan solos. Pon descuento si aplica."),
 ("4. Pagos", "Cada Yape, Plin o transferencia, con N° de operación y el ID de venta que paga. Así la deuda se actualiza sola."),
 ("5. Clases", "Cada clase dictada: alumno, horas, tema y link de la grabación. Descuenta horas del alumno y calcula el pago al tutor. 'Reprogramada' no descuenta."),
 ("6. Avance", "Notas de simulacros y prácticas. Sirve para reportes a padres y para mostrar progreso."),
 ("7. Tutores", "Hoy solo tú (pago 0). Cuando contrates, agrega su tarifa por hora y el sistema calcula cuánto le debes."),
 ("Catálogo", "Copia de tu tarifario aprobado (archivo precios/tarifario.csv, la misma fuente que usa la web). No lo edites aquí: cambia el tarifario y pídeme sincronizar, o escríbelo en el hilo de precios. Los códigos de producto son los mismos en la web y en el SIG."),
 ("Dashboard", "Elige el mes en la celda amarilla y verás ingresos, horas, leads, conversión y resultados por campaña."),
 ("Códigos de campaña", "SPRINT-ULIMA, SPRINT-PUCP, REFUERZO-CEPU, MEDICINA-UPCH, VERANO-PUCP, VERANO-ULIMA, CACHIMBOS, REFERIDOS, ORGANICO. Se editan en la hoja Listas."),
 ("Control", "Revisión automática de errores (IDs que no existen, pagos duplicados, horas negativas, paquetes vencidos). Si algo dice 'Revisar', corrígelo. El Dashboard muestra el total de alertas."),
 ("Cambios y respaldos", "La hoja Cambios guarda quién cambió qué y cuándo (requiere instalar el script ASTROS_SIG_script.gs, ver guía). El mismo script crea una copia de respaldo semanal en tu Drive. Además, Google Sheets guarda todo en Archivo > Historial de versiones."),
 ("Desde el celular", "Usa la app de Google Sheets. Toca una celda con flecha para elegir de la lista. Para fechas escribe dd/mm/aaaa."),
]
for i, (a, b) in enumerate(rows, 4):
    ini[f"B{i}"] = a; ini[f"B{i}"].font = Font(bold=True, color=NAVY)
    ini[f"C{i}"] = b; ini[f"C{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ini.row_dimensions[i].height = 36

del wb["Sheet"]
order = ["Inicio", "Dashboard", "Control", "Leads", "Alumnos", "Ventas", "Pagos", "Clases", "Avance", "Tutores", "Catalogo", "Cambios", "Listas"]
wb._sheets = [wb[n] for n in order]

if SAMPLE:
    import datetime as dt
    d = dt.date.today()
    ld["B2"], ld["C2"], ld["F2"], ld["H2"], ld["I2"], ld["J2"] = d, "Ana Test", "PUCP", "Instagram", "SPRINT-PUCP", "Inscrito"
    ld["B3"], ld["C3"], ld["F3"], ld["H3"], ld["I3"], ld["J3"], ld["K3"] = d, "Luis Test", "ULIMA", "WhatsApp", "SPRINT-ULIMA", "Contactado", d - dt.timedelta(days=3)
    al["B2"], al["M2"] = "Ana Test", "Activo"
    ve["B2"], ve["C2"], ve["E2"], ve["H2"], ve["M2"] = d, "A001", "pre-pack12", 40, "SPRINT-PUCP"
    pg["A2"], pg["C2"], pg["D2"], pg["E2"] = d, "A001", "V001", 300
    cl["A2"], cl["C2"], cl["F2"], cl["I2"], cl["J2"] = d, "A001", "Martini", 2, "Asistió"
    cl["A3"], cl["C3"], cl["F3"], cl["I3"], cl["J3"] = d, "A001", "Martini", 1.5, "Reprogramada"
wb.save(OUT)
print("ok", OUT)

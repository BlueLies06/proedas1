"""Valida preguntas.csv, guarda respaldo y log si cambió, y genera ASTROS_Banco_Preguntas.xlsx.
Si hay errores, los anota en logs/errores.log y NO regenera el Excel.
Uso: python3 validar_banco.py [--restaurar backups/AAAAMMDD-HHMMSS]"""
import csv, datetime as dt, hashlib, os, shutil, sys
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV = os.path.join(BASE, "preguntas.csv")
COLS = ["id", "universidad", "examen", "area", "curso", "tema", "dificultad", "texto", "enunciado", "A", "B", "C", "D", "E", "respuesta", "solucion", "tiempo_seg", "fuente", "estado"]
UNIS = {"PUCP", "ULIMA", "UPC", "UPCH", "Universitario", "Escolar"}

def validar(filas):
    err, ids = [], set()
    for n, r in enumerate(filas, 2):
        p = f"fila {n} ({r.get('id') or 'sin id'})"
        for c in ("id", "universidad", "area", "enunciado", "A", "B", "C", "D", "E", "respuesta", "solucion"):
            if not (r.get(c) or "").strip(): err.append(f"{p}: falta '{c}'")
        if r.get("id") in ids: err.append(f"{p}: id repetido")
        ids.add(r.get("id"))
        if r.get("universidad") not in UNIS: err.append(f"{p}: universidad '{r.get('universidad')}' no válida ({', '.join(sorted(UNIS))})")
        if r.get("respuesta") not in list("ABCDE"): err.append(f"{p}: respuesta debe ser A, B, C, D o E")
        ops = [(r.get(k) or "").strip() for k in "ABCDE"]
        if all(ops) and len(set(ops)) < 5: err.append(f"{p}: hay alternativas repetidas")
        if r.get("dificultad") not in ("1", "2", "3"): err.append(f"{p}: dificultad debe ser 1, 2 o 3")
        if not (r.get("tiempo_seg") or "").isdigit(): err.append(f"{p}: tiempo_seg debe ser un número")
    return err

def leer(ruta=CSV):
    with open(ruta, encoding="utf-8") as f:
        rd = csv.DictReader(f)
        if rd.fieldnames != COLS: return None, [f"columnas incorrectas: se esperaba {','.join(COLS)}"]
        return list(rd), []

def construir_xlsx(filas, salida):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    NAVY, GOLD, TEAL = "0B2533", "F9C74F", "1FA2A6"
    wb = Workbook(); ini = wb.active; ini.title = "Inicio"
    ini.column_dimensions["A"].width = 24; ini.column_dimensions["B"].width = 90
    ini["A1"] = "ASTROS · Banco de preguntas"; ini["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ini["A1"].fill = ini["B1"].fill = PatternFill("solid", fgColor=NAVY)
    info = [("Cómo se usa", "Filtra la hoja Preguntas por universidad, área, tema o dificultad (flechas del encabezado). La hoja Resumen cuenta cuántas preguntas tienes de cada tipo."),
            ("Agregar preguntas", "La fuente es preguntas.csv. Agrega filas ahí (o pídeselas a Claude) y corre herramientas/validar_banco.py: revisa errores, respalda y regenera este Excel."),
            ("Simulacros", "herramientas/generar_simulacro.py arma un PDF listo para enviar al alumno, con hoja de claves y solucionario aparte."),
            ("Resultados", "Registra los puntajes de cada simulacro en la hoja Avance del SIG."),
            ("Formato del examen", "Las áreas por universidad están en la hoja Estructura. Confirma siempre con la guía oficial del postulante: los formatos cambian.")]
    for i, (a, b) in enumerate(info, 3):
        ini[f"A{i}"] = a; ini[f"A{i}"].font = Font(bold=True, color=NAVY); ini[f"B{i}"] = b
        ini[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top"); ini.row_dimensions[i].height = 34
    ws = wb.create_sheet("Preguntas")
    widths = [10, 10, 18, 20, 18, 20, 8, 50, 50, 22, 22, 22, 22, 22, 9, 50, 9, 22, 16]
    for j, (c, w) in enumerate(zip(COLS, widths), 1):
        ws.cell(1, j, c).font = Font(bold=True, color=NAVY); ws.cell(1, j).fill = PatternFill("solid", fgColor=GOLD)
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(filas, 2):
        for j, c in enumerate(COLS, 1):
            v = r[c]; ws.cell(i, j, int(v) if c in ("dificultad", "tiempo_seg") else v)
            ws.cell(i, j).alignment = Alignment(wrap_text=c in ("texto", "enunciado", "solucion"), vertical="top")
    ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(filas)+1}"
    es = wb.create_sheet("Estructura")
    est = [("Universidad", "Áreas a preparar (referencial)", "Confianza"),
           ("PUCP", "Razonamiento matemático, matemática básica, comprensión lectora, redacción", "Confirmar en la guía del postulante PUCP"),
           ("ULIMA", "Razonamiento matemático, razonamiento verbal, comprensión lectora", "Confirmar en ulima.edu.pe"),
           ("UPC", "Evaluación integral online: aptitud matemática y verbal", "Confirmar al inscribirse (610-5030)"),
           ("UPCH", "Matemática, razonamiento verbal, biología, química, física (énfasis ciencias de la salud)", "Confirmar en cayetano.edu.pe"),
           ("Universitario", "Prácticas y exámenes por curso (Cálculo, Física, Química, Estadística...)", "Según el sílabo del curso")]
    for i, row in enumerate(est, 1):
        for j, v in enumerate(row, 1):
            es.cell(i, j, v); es.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
            if i == 1: es.cell(i, j).font = Font(bold=True, color=NAVY); es.cell(i, j).fill = PatternFill("solid", fgColor=GOLD)
    for c, w in zip("ABC", (14, 70, 40)): es.column_dimensions[c].width = w
    rs = wb.create_sheet("Resumen")
    rs["A1"], rs["B1"], rs["C1"], rs["D1"], rs["E1"], rs["F1"] = "Universidad", "Área", "Fácil (1)", "Media (2)", "Difícil (3)", "Total"
    for c in "ABCDEF": rs[f"{c}1"].font = Font(bold=True, color="FFFFFF"); rs[f"{c}1"].fill = PatternFill("solid", fgColor=TEAL)
    combos = sorted({(r["universidad"], r["area"]) for r in filas})
    for i, (u, a) in enumerate(combos, 2):
        rs[f"A{i}"], rs[f"B{i}"] = u, a
        for k, col in enumerate("CDE", 1):
            rs[f"{col}{i}"] = f'=COUNTIFS(Preguntas!$B:$B,A{i},Preguntas!$D:$D,B{i},Preguntas!$G:$G,{k})'
        rs[f"F{i}"] = f"=SUM(C{i}:E{i})"
    t = len(combos) + 2
    rs[f"A{t}"] = "TOTAL"; rs[f"A{t}"].font = Font(bold=True)
    for col in "CDEF": rs[f"{col}{t}"] = f"=SUM({col}2:{col}{t-1})"
    rs.column_dimensions["A"].width = 14; rs.column_dimensions["B"].width = 26
    wb.save(salida)

def exportar_json(filas, salida):
    """preguntas.json para la web: agrupado por universidad, mismo contenido que preguntas.csv."""
    import json
    out = {"version": dt.datetime.now().strftime("%Y-%m-%dT%H:%M:%S"), "universidades": {}}
    for r in filas:
        out["universidades"].setdefault(r["universidad"], []).append({
            "id": r["id"], "examen": r["examen"], "area": r["area"], "curso": r["curso"], "tema": r["tema"],
            "dificultad": int(r["dificultad"]), "texto": r["texto"] or None, "enunciado": r["enunciado"],
            "opciones": {k: r[k] for k in "ABCDE"}, "respuesta": r["respuesta"], "solucion": r["solucion"],
            "tiempo_seg": int(r["tiempo_seg"])})
    tmp = salida + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f: json.dump(out, f, ensure_ascii=False, indent=1)
    os.replace(tmp, salida)

def exportar_web(filas, salida, n=20):
    """20 preguntas por universidad en el formato del test de la web (correcta = índice 0-based).
    Las preguntas de un mismo texto van completas; el texto viene en 'texto'."""
    import json
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from generar_simulacro import seleccionar
    out = {"actualizado": dt.date.today().isoformat(), "universidades": {}}
    for u in ("PUCP", "ULIMA", "UPC", "UPCH"):
        if not any(r["universidad"] == u for r in filas): continue
        out["universidades"][u] = [{
            "id": r["id"].lower(), "area": r["area"], "tema": r["tema"], "dificultad": int(r["dificultad"]),
            "texto": r["texto"] or None, "enunciado": r["enunciado"], "opciones": [r[k] for k in "ABCDE"],
            "correcta": "ABCDE".index(r["respuesta"]), "solucion": r["solucion"], "fuente": "propia"}
            for r in seleccionar(filas, u, None, n, 1)]
    os.makedirs(os.path.dirname(salida), exist_ok=True)
    tmp = salida + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f: json.dump(out, f, ensure_ascii=False, indent=1)
    os.replace(tmp, salida)

def huella(ruta):
    return hashlib.sha256(open(ruta, "rb").read()).hexdigest()

def main():
    ahora = dt.datetime.now(); os.makedirs(os.path.join(BASE, "logs"), exist_ok=True)
    def log(nombre, texto):
        with open(os.path.join(BASE, "logs", nombre), "a", encoding="utf-8") as f: f.write(f"{ahora:%Y-%m-%d %H:%M:%S} {texto}\n")
    if len(sys.argv) == 3 and sys.argv[1] == "--restaurar":
        origen = os.path.join(BASE, sys.argv[2], "preguntas.csv") if not os.path.isabs(sys.argv[2]) else os.path.join(sys.argv[2], "preguntas.csv")
        d = os.path.join(BASE, "backups", f"{ahora:%Y%m%d-%H%M%S}-antes-de-restaurar"); os.makedirs(d); shutil.copy2(CSV, d)
        shutil.copy2(origen, CSV); log("cambios.log", f"restaurado desde {sys.argv[2]} (versión previa en {d})")
    filas, err = leer()
    if not err: err = validar(filas)
    if err:
        log("errores.log", f"{len(err)} error(es):\n    " + "\n    ".join(err))
        print("ERRORES (no se regeneró el Excel):\n  " + "\n  ".join(err)); sys.exit(1)
    bdir = os.path.join(BASE, "backups"); os.makedirs(bdir, exist_ok=True)
    previos = sorted(d for d in os.listdir(bdir) if os.path.exists(os.path.join(bdir, d, "preguntas.csv")))
    ultimo = os.path.join(bdir, previos[-1], "preguntas.csv") if previos else None
    if not ultimo or huella(ultimo) != huella(CSV):
        antes = {r["id"]: r for r in csv.DictReader(open(ultimo, encoding="utf-8"))} if ultimo else {}
        ahora_ids = {r["id"]: r for r in filas}
        nuevas = sorted(set(ahora_ids) - set(antes)); quitadas = sorted(set(antes) - set(ahora_ids))
        editadas = sorted(i for i in set(antes) & set(ahora_ids) if antes[i] != ahora_ids[i])
        d = os.path.join(bdir, f"{ahora:%Y%m%d-%H%M%S}"); os.makedirs(d, exist_ok=True); shutil.copy2(CSV, d)
        log("cambios.log", f"respaldo {os.path.relpath(d, BASE)} · {len(filas)} preguntas · nuevas {nuevas or '-'} · editadas {editadas or '-'} · quitadas {quitadas or '-'}")
        for viejo in sorted(os.listdir(bdir))[:-60]: shutil.rmtree(os.path.join(bdir, viejo), ignore_errors=True)
    construir_xlsx(filas, os.path.join(BASE, "ASTROS_Banco_Preguntas.xlsx"))
    exportar_json(filas, os.path.join(BASE, "preguntas.json"))
    exportar_web(filas, os.path.join(os.path.dirname(BASE), "sig", "preguntas-web.json"))
    print(f"OK: {len(filas)} preguntas válidas. Excel actualizado.")

if __name__ == "__main__":
    main()
